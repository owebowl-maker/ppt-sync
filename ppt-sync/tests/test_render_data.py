# -*- coding: utf-8 -*-
"""阶段 C：跨 run 数字替换、图表更新、缺 key 报错、其余内容不变。"""
import pandas as pd
import pytest
from pptx import Presentation

from ppt_sync.render_data import (
    RenderLog, load_data, replace_numbers_in_paragraph, render_all,
)


def _slide0(path):
    return Presentation(str(path)).slides[0]


def _para_with_runs(tf, texts):
    p = tf.paragraphs[0]
    p.text = ""
    for t in texts:
        r = p.add_run()
        r.text = t
    return p


def test_cross_run_replace(tmp_path):
    """'128' 与 ' 家' 拆在两个 run：替换数字、保留单位与格式。"""
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    tb = s.shapes.add_textbox(0, 0, 100, 100)
    p = _para_with_runs(tb.text_frame, ["累计合作案例 ", "128", " 家"])
    # 格式标记：给第二个 run 加粗，确认格式随 run 保留
    from pptx.util import Pt
    p.runs[1].font.bold = True
    log = RenderLog()
    n = replace_numbers_in_paragraph(p, 156, "{:,.0f}", "t", log)
    assert n == 1
    assert p.text == "累计合作案例 156 家"
    assert p.runs[1].font.bold           # 数字 run 的格式保留
    assert p.runs[2].text == " 家"        # 单位 run 未动


def test_keep_suffix_and_decimal(tmp_path):
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    tb = s.shapes.add_textbox(0, 0, 100, 100)
    p = _para_with_runs(tb.text_frame, ["销售额 ", "3.2", "亿元"])
    log = RenderLog()
    replace_numbers_in_paragraph(p, 3.8, "{:.1f}", "t", log)
    assert p.text == "销售额 3.8亿元"


def test_render_full_pipeline(samples, tagged, tmp_path_factory):
    tagged = tagged[0]
    """对打标样例跑 render：数值全部同步、单位保留、图表更新、其余不变。"""
    out = tmp_path_factory.mktemp("rendered")
    data = tmp_path_factory.mktemp("data") / "data.xlsx"
    metrics = pd.DataFrame([
        {"metric_key": "case_count", "值": 156, "单位": "个", "显示格式": "{:,.0f}"},
        {"metric_key": "sales_total", "值": 3.8, "单位": "亿元", "显示格式": "{:.1f}"},
        {"metric_key": "partner_count", "值": 52, "单位": "家", "显示格式": "{:,.0f}"},
    ])
    chart = pd.DataFrame({"季度": ["2025Q3", "2025Q4", "2026Q1", "2026Q2"],
                          "销售额(亿元)": [2.1, 2.4, 2.8, 3.8]})
    with pd.ExcelWriter(data, engine="xlsxwriter") as w:
        metrics.to_excel(w, sheet_name="metrics", index=False)
        chart.to_excel(w, sheet_name="chart_sales_total", index=False)
    written, log, _ = render_all(tagged, data, out, report_period="2026Q3")
    assert len(written) == 3
    assert not log.errors

    prs = Presentation(str(out / "PPT-1_2026Q3.pptx"))

    def texts_of(slide):
        return {sh.text_frame.text for sh in slide.shapes if sh.has_text_frame}

    assert "销售额 3.8亿元" in texts_of(prs.slides[2])
    assert "累计合作厂家 52家" in texts_of(prs.slides[3])
    assert any(t.endswith("156 个成功落地") for t in texts_of(prs.slides[4]))
    chart_vals = [sh.chart.plots[0].series[0].values for sh in prs.slides[2].shapes if sh.has_chart]
    assert list(chart_vals[0])[-1] == pytest.approx(3.8)

    # 未打标页逐字符不变
    orig = Presentation(str(tagged / "PPT-1.pptx"))
    for i, (sa, sb) in enumerate(zip(orig.slides, prs.slides), 1):
        if i in (3, 4, 5):
            continue
        ta = [sh.text_frame.text for sh in sa.shapes if sh.has_text_frame]
        tb = [sh.text_frame.text for sh in sb.shapes if sh.has_text_frame]
        assert ta == tb


def test_missing_metric_key_errors(samples, tagged, tmp_path_factory):
    tagged = tagged[0]
    """PPT 有标签但数据源缺 key → 必须报错退出，不能静默跳过。"""
    out = tmp_path_factory.mktemp("rendered2")
    data = tmp_path_factory.mktemp("data2") / "data.xlsx"
    pd.DataFrame([{"metric_key": "case_count", "值": 156, "单位": "个",
                   "显示格式": "{:,.0f}"}]).to_excel(data, index=False, sheet_name="metrics")
    written, log, _ = render_all(tagged, data, out)
    assert any("sales_total" in e for e in log.errors)
    assert any("partner_count" in e for e in log.errors)


def test_unused_metric_warning(samples, tagged, tmp_path_factory):
    tagged = tagged[0]
    """数据源有、PPT 无标签的 key → 记入 warnings。"""
    out = tmp_path_factory.mktemp("rendered3")
    data = tmp_path_factory.mktemp("data3") / "data.xlsx"
    metrics = pd.DataFrame([
        {"metric_key": "case_count", "值": 156, "单位": "个", "显示格式": "{:,.0f}"},
        {"metric_key": "sales_total", "值": 3.8, "单位": "亿元", "显示格式": "{:.1f}"},
        {"metric_key": "partner_count", "值": 52, "单位": "家", "显示格式": "{:,.0f}"},
        {"metric_key": "ghost_metric", "值": 1, "单位": "", "显示格式": ""},
    ])
    with pd.ExcelWriter(data, engine="xlsxwriter") as w:
        metrics.to_excel(w, sheet_name="metrics", index=False)
    written, log, _ = render_all(tagged, data, out)
    assert any("ghost_metric" in wmsg for wmsg in log.warnings)
