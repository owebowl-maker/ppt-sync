# -*- coding: utf-8 -*-
"""阶段 B：打标读写往返、XML 仅标签两处变化、registry。"""
import re
import zipfile

import pandas as pd
import pytest
from lxml import etree
from pptx import Presentation

from conftest import CONFIG, KNOWN_PAIRS


@pytest.fixture(scope="module")
def tagged(samples, index, tmp_path_factory):
    from ppt_sync.tag import tag_all, save_registry
    out = tmp_path_factory.mktemp("tagged")
    # 人工确认表：按 KNOWN_PAIRS 的分组填
    rows = [
        {"group_id": 0, "成员": "PPT-1.pptx#3, PPT-2.pptx#2",
         "确认(Y/N)": "Y", "指标key(英文)": "sales_total", "指标名称": "销售额"},
        {"group_id": 1, "成员": "PPT-1.pptx#4, PPT-2.pptx#3",
         "确认(Y/N)": "Y", "指标key(英文)": "partner_count", "指标名称": "合作厂家数"},
        {"group_id": 2, "成员": "PPT-1.pptx#5, PPT-3.pptx#13",
         "确认(Y/N)": "Y", "指标key(英文)": "case_count", "指标名称": "合作案例数"},
    ]
    mapping = tmp_path_factory.mktemp("conf") / "confirmed.xlsx"
    pd.DataFrame(rows).to_excel(mapping, index=False)
    written, registry, report = tag_all(samples, mapping, out, index)
    return out, registry, report


def test_tag_roundtrip(tagged):
    """打标后重新读取，标签仍在，且与确认表一致。"""
    out, registry, report = tagged
    prs = Presentation(str(out / "PPT-1.pptx"))
    assert "metric=sales_total" in (prs.slides[2].name or "")
    assert "metric=case_count" in (prs.slides[4].name or "")
    names = {sh.name for sh in prs.slides[4].shapes}
    assert any("metric:case_count" in n for n in names)


def test_tag_by_slide_id(tagged):
    out, registry, report = tagged
    assert report.slides_tagged == 6
    # registry 的页码与 KNOWN_PAIRS 一致
    reg = pd.DataFrame(registry)
    for deck, no in KNOWN_PAIRS:
        assert no in set(reg[reg["deck"] == deck]["当前页码"])


def test_xml_diff_only_names(samples, tagged):
    """打标前后 slide XML 只允许 p:cSld/@name 与 cNvPr/@name 两处差异。"""
    out, _, _ = tagged
    for deck in ("PPT-1.pptx", "PPT-2.pptx", "PPT-3.pptx"):
        with zipfile.ZipFile(samples / deck) as za, \
             zipfile.ZipFile(out / deck) as zb:
            na = {n for n in za.namelist() if re.match(r"ppt/slides/slide\d+\.xml$", n)}
            assert na == {n for n in zb.namelist() if re.match(r"ppt/slides/slide\d+\.xml$", n)}
            for n in na:
                da, db = etree.fromstring(za.read(n)), etree.fromstring(zb.read(n))
                for ea, eb in zip(da.iter(), db.iter()):
                    assert ea.tag == eb.tag
                    diff_keys = {k for k in set(ea.attrib) ^ set(eb.attrib)}
                    diff_vals = {k for k in set(ea.attrib) & set(eb.attrib)
                                 if ea.attrib[k] != eb.attrib[k]}
                    changed = diff_keys | diff_vals
                    for k in changed:
                        assert k == "name", f"{deck}/{n}: 意外属性变化 {k}"
                    assert (ea.text or "") == (eb.text or ""), f"{deck}/{n}: 文本被改动"


def test_shape_name_parse():
    from ppt_sync.tag import _parse_shape_name, shape_metric_key
    assert _parse_shape_name("原名|metric:a") == ("原名", ["a"])
    assert _parse_shape_name("原名|metric:a|metric:b") == ("原名", ["a", "b"])
    assert shape_metric_key(type("S", (), {"name": "x|metric:case_count"})()) == "case_count"
    assert shape_metric_key(type("S", (), {"name": "TextBox 1"})()) is None
